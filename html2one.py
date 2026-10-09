#!/usr/bin/env python3
"""html2one — ko'p faylli HTML loyihani (zip yoki papka) BITTA faylga o'tkazadi.

Natijalar (hammasi bitta-bitta fayl):
  <nom>.pptx  — har bir slayd brauzerda qanday ko'rinsa, aynan shunday (rasm sifatida)
  <nom>.pdf   — xuddi shu sahifalar, PDF ko'rinishida
  <nom>.html  — barcha CSS/JS/rasm/shriftlar ichiga joylangan, interaktiv bitta HTML

Ishlatish:
  pip install playwright python-pptx pillow beautifulsoup4
  python html2one.py                          # ~/Downloads/anor-delivery.zip ni oladi
  python html2one.py boshqa-fayl.zip          # yoki istalgan zip / papka / index.html

Brauzer sifatida kompyuterdagi Chrome/Edge ishlatiladi. Ular bo'lmasa, bir marta:
  python -m playwright install chromium

Qo'shimcha:
  --out DIR            natijalar papkasi (standart: zip yonida)
  --size 1920x1080     brauzer oynasi o'lchami (slayd o'lchami shundan olinadi)
  --mode auto|reveal|sections|keys|page
                       slaydlarni qanday ajratish (standart: auto — o'zi aniqlaydi)
  --selector CSS       --mode sections uchun slayd elementlari selektori
  --wait MS            har slayddan keyin animatsiya kutish vaqti (standart 800)
  --scale N            rasm aniqligi (standart 2 = retina)
  --only pptx,pdf,html kerakli formatlar
  --no-remote          internetdagi resurslarni (Google Fonts va h.k.) yuklamaslik
"""
from __future__ import annotations

import argparse
import base64
import functools
import http.server
import io
import json
import mimetypes
import os
import re
import shutil
import sys
import tempfile
import threading
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

mimetypes.add_type("font/woff2", ".woff2")
mimetypes.add_type("font/woff", ".woff")
mimetypes.add_type("font/ttf", ".ttf")
mimetypes.add_type("font/otf", ".otf")
mimetypes.add_type("image/svg+xml", ".svg")
mimetypes.add_type("image/webp", ".webp")
mimetypes.add_type("image/avif", ".avif")
mimetypes.add_type("application/json", ".json")
mimetypes.add_type("text/javascript", ".mjs")


def log(msg: str) -> None:
    print(msg, flush=True)


# ---------------------------------------------------------------- kirish

def prepare_source(src: Path, work: Path) -> Path:
    """Zipni ochadi va index.html joylashgan papkani qaytaradi."""
    if src.is_dir():
        root = src
    elif zipfile.is_zipfile(src):
        root = work / "src"
        with zipfile.ZipFile(src) as z:
            for info in z.infolist():
                name = info.filename
                # macOS zip ichidagi keraksiz fayllar va xavfli yo'llarni o'tkazib yuboramiz
                if name.startswith("__MACOSX/") or "/._" in name or name.startswith("._"):
                    continue
                target = (root / name).resolve()
                if not str(target).startswith(str(root.resolve())):
                    continue
                if info.is_dir():
                    target.mkdir(parents=True, exist_ok=True)
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                with z.open(info) as fsrc, open(target, "wb") as fdst:
                    shutil.copyfileobj(fsrc, fdst)
    elif src.suffix.lower() in (".html", ".htm"):
        return src.parent
    else:
        sys.exit(f"Xato: {src} zip, papka yoki .html emas")

    candidates = sorted(root.rglob("index.html"), key=lambda p: len(p.parts))
    if not candidates:
        candidates = sorted(root.rglob("*.html"), key=lambda p: len(p.parts))
    if not candidates:
        sys.exit("Xato: ichida .html fayl topilmadi")
    return candidates[0].parent


def entry_file(root: Path, src: Path) -> Path:
    if src.suffix.lower() in (".html", ".htm"):
        return src
    idx = root / "index.html"
    return idx if idx.exists() else sorted(root.glob("*.html"))[0]


class QuietHandler(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass


def serve(root: Path) -> tuple[http.server.ThreadingHTTPServer, int]:
    handler = functools.partial(QuietHandler, directory=str(root))
    httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd, httpd.server_address[1]


# ---------------------------------------------------------------- suratga olish

SETTLE_JS = """async () => {
  try { await document.fonts.ready; } catch (e) {}
  const imgs = [...document.images].filter(i => !i.complete);
  await Promise.race([
    Promise.all(imgs.map(i => new Promise(r => { i.onload = i.onerror = r; }))),
    new Promise(r => setTimeout(r, 4000)),
  ]);
}"""

DETECT_JS = """() => {
  const vw = innerWidth, vh = innerHeight;
  const out = {
    reveal: !!(window.Reveal && typeof Reveal.getTotalSlides === 'function'),
    scrollH: document.documentElement.scrollHeight,
    vh, vw, sections: null,
  };
  const sels = ['section.slide', '.slide', '[data-slide]', '.swiper-slide',
                '.page', 'main > section', 'body > section', 'section'];
  for (const sel of sels) {
    const els = [...document.querySelectorAll(sel)].filter(e => !e.parentElement.closest(sel));
    if (els.length < 2) continue;
    const vis = els.filter(e => {
      const r = e.getBoundingClientRect(), cs = getComputedStyle(e);
      return r.width > 0 && r.height > 0 && cs.display !== 'none' && cs.visibility !== 'hidden'
             && parseFloat(cs.opacity) > 0.05;
    });
    const big = vis.filter(e => {
      const r = e.getBoundingClientRect();
      return r.height >= vh * 0.6 && r.width >= vw * 0.6;
    });
    out.sections = { sel, total: els.length, visible: vis.length, big: big.length };
    break;
  }
  return out;
}"""


def launch_browser(p):
    """Playwright Chromium bo'lmasa — kompyuterdagi Chrome yoki Edge'dan foydalanadi."""
    exe = os.environ.get("HTML2ONE_CHROMIUM")
    attempts = [{"executable_path": exe}] if exe else []
    attempts += [{}, {"channel": "chrome"}, {"channel": "msedge"}]
    last = None
    for opts in attempts:
        try:
            return p.chromium.launch(**opts)
        except Exception as e:
            last = e
    sys.exit("Xato: brauzer topilmadi. Chrome o'rnating yoki: python -m playwright install chromium\n"
             f"({last})")


def shot(page, scale_jpeg: bool, clip=None, full=False) -> bytes:
    kw = {"type": "jpeg", "quality": 92} if scale_jpeg else {"type": "png"}
    if clip:
        kw["clip"] = clip
    if full:
        kw["full_page"] = True
    return page.screenshot(animations="disabled", caret="hide", **kw)


def capture(entry_url: str, args) -> list[dict]:
    from playwright.sync_api import sync_playwright

    w, h = args.width, args.height
    slides: list[dict] = []
    with sync_playwright() as p:
        browser = launch_browser(p)
        ctx = browser.new_context(viewport={"width": w, "height": h},
                                  device_scale_factor=args.scale)
        page = ctx.new_page()
        page.goto(entry_url, wait_until="load", timeout=60000)
        try:
            page.wait_for_load_state("networkidle", timeout=8000)
        except Exception:
            pass
        page.evaluate(SETTLE_JS)
        page.wait_for_timeout(args.wait)

        info = page.evaluate(DETECT_JS)
        mode = args.mode
        if mode == "auto":
            sec = info["sections"]
            if info["reveal"]:
                mode = "reveal"
            elif (sec and sec["big"] >= 2 and sec["visible"] == sec["total"]
                  and sec["big"] >= 0.6 * sec["total"]):
                mode = "sections"
                args.selector = args.selector or sec["sel"]
            elif info["scrollH"] <= info["vh"] * 1.15:
                mode = "keys"
            else:
                mode = "page"
        log(f"  rejim: {mode}")

        jpeg = not args.png
        if mode == "reveal":
            page.evaluate("""() => { Reveal.configure({ fragments: false, transition: 'none',
                              backgroundTransition: 'none', controls: false, progress: false });
                              Reveal.slide(0, 0); }""")
            for _ in range(1000):
                page.wait_for_timeout(args.wait)
                page.evaluate(SETTLE_JS)
                text = page.evaluate("() => Reveal.getCurrentSlide()?.innerText || ''")
                slides.append({"img": shot(page, jpeg), "text": text})
                if page.evaluate("() => Reveal.isLastSlide()"):
                    break
                page.evaluate("() => Reveal.next()")

        elif mode == "sections":
            sel = args.selector or "section"
            count = page.evaluate(
                "(s) => [...document.querySelectorAll(s)].filter(e => !e.parentElement.closest(s)).length", sel)
            for i in range(count):
                handle = page.evaluate_handle(
                    "([s, i]) => [...document.querySelectorAll(s)].filter(e => !e.parentElement.closest(s))[i]",
                    [sel, i]).as_element()
                handle.scroll_into_view_if_needed()
                page.wait_for_timeout(args.wait)
                page.evaluate(SETTLE_JS)
                text = handle.evaluate("e => e.innerText || ''")
                img = handle.screenshot(animations="disabled", caret="hide",
                                        **({"type": "jpeg", "quality": 92} if jpeg else {"type": "png"}))
                slides.append({"img": img, "text": text})

        elif mode == "keys":
            prev = None
            same = 0
            for _ in range(500):
                img = shot(page, jpeg)
                if prev is not None and looks_same(img, prev):
                    same += 1
                    if same >= 2:   # 2 marta bosildi — o'zgarmadi, demak oxirgi slayd
                        break
                else:
                    same = 0
                    slides.append({"img": img, "text": ""})
                    prev = img
                page.keyboard.press("ArrowRight" if same == 0 else "PageDown")
                page.wait_for_timeout(args.wait)
                page.evaluate(SETTLE_JS)

        else:  # page — uzun sahifa: oxirigacha aylantirib, keyin bo'laklarga bo'lamiz
            page.evaluate("""async () => {
                for (let y = 0; y < document.documentElement.scrollHeight; y += innerHeight / 2) {
                    scrollTo(0, y); await new Promise(r => setTimeout(r, 120)); }
                scrollTo(0, 0); }""")
            page.wait_for_timeout(args.wait)
            page.evaluate(SETTLE_JS)
            full = page.screenshot(full_page=True, type="png", animations="disabled", caret="hide")
            slides.extend(slice_long(full, w, h, args.scale, jpeg))

        browser.close()

    fitted = []
    for s in slides:
        for img in fit_canvas(s["img"], w, h, args.scale, jpeg):
            fitted.append({"img": img, "text": s["text"]})
    return fitted


def looks_same(a: bytes, b: bytes) -> bool:
    """Ikki skrinshot deyarli bir xilmi (kichik animatsiya/soat farqini hisobga olmaydi)."""
    from PIL import Image, ImageChops, ImageStat

    ta = Image.open(io.BytesIO(a)).convert("L").resize((96, 54))
    tb = Image.open(io.BytesIO(b)).convert("L").resize((96, 54))
    return ImageStat.Stat(ImageChops.difference(ta, tb)).mean[0] < 1.5


def fit_canvas(img: bytes, w: int, h: int, scale: float, jpeg: bool) -> list[bytes]:
    """Rasmni slayd o'lchamiga keltiradi: past bo'lsa to'ldiradi, baland bo'lsa bo'laklaydi."""
    from PIL import Image

    im = Image.open(io.BytesIO(img)).convert("RGB")
    cw, ch = round(w * scale), round(h * scale)
    if im.size == (cw, ch):
        return [img]
    if im.width != cw:
        im = im.resize((cw, max(1, round(im.height * cw / im.width))), Image.LANCZOS)
    if im.height > ch * 1.02:
        buf = io.BytesIO()
        im.save(buf, "PNG")
        return [x["img"] for x in slice_long(buf.getvalue(), w, h, scale, jpeg)]
    bg = im.getpixel((im.width // 2, im.height - 1))
    canvas = Image.new("RGB", (cw, ch), bg)
    canvas.paste(im.crop((0, 0, cw, min(im.height, ch))), (0, 0))
    buf = io.BytesIO()
    canvas.save(buf, "JPEG" if jpeg else "PNG", **({"quality": 92} if jpeg else {}))
    return [buf.getvalue()]


def slice_long(png: bytes, w: int, h: int, scale: float, jpeg: bool) -> list[dict]:
    from PIL import Image

    im = Image.open(io.BytesIO(png)).convert("RGB")
    step = int(h * scale)
    out = []
    for top in range(0, im.height, step):
        part = im.crop((0, top, im.width, min(top + step, im.height)))
        if part.height < step:
            # oxirgi bo'lak qisqa bo'lsa — sahifa pastki rangi bilan to'ldiramiz
            bg = part.getpixel((part.width // 2, part.height - 1))
            canvas = Image.new("RGB", (im.width, step), bg)
            canvas.paste(part, (0, 0))
            part = canvas
        buf = io.BytesIO()
        part.save(buf, "JPEG" if jpeg else "PNG", **({"quality": 92} if jpeg else {}))
        out.append({"img": buf.getvalue(), "text": ""})
    return out


# ---------------------------------------------------------------- PPTX / PDF

def write_pptx(slides: list[dict], path: Path, w: int, h: int) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    emu_w = Emu(12192000)                       # 13.333" — standart 16:9 kenglik
    emu_h = Emu(int(12192000 * h / w))
    prs.slide_width, prs.slide_height = emu_w, emu_h
    blank = prs.slide_layouts[6]
    for s in slides:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(s["img"]), 0, 0, width=emu_w, height=emu_h)
        if s.get("text", "").strip():
            slide.notes_slide.notes_text_frame.text = s["text"].strip()
    prs.save(path)


def write_pdf(slides: list[dict], path: Path, w: int) -> None:
    from PIL import Image

    pages = [Image.open(io.BytesIO(s["img"])).convert("RGB") for s in slides]
    dpi = pages[0].width / 13.333
    pages[0].save(path, "PDF", save_all=True, append_images=pages[1:], resolution=dpi, quality=95)


# ---------------------------------------------------------------- bitta HTML

URL_RE = re.compile(r"url\(\s*(['\"]?)([^'\")]+)\1\s*\)", re.I)
IMPORT_RE = re.compile(
    r"@import\s+(?:url\(\s*['\"]?([^'\")]+)['\"]?\s*\)|['\"]([^'\"]+)['\"])\s*([^;]*);", re.I)
FONTFACE_RE = re.compile(r"@font-face\s*\{[^}]*\}", re.I)
FONT_RANK = [(".woff2", "woff2"), (".woff", "woff"), (".ttf", "truetype"), (".otf", "opentype")]


def prune_font_face(block: str) -> str:
    """@font-face ichida bir nechta format bo'lsa — faqat eng yaxshisini qoldiradi (eot/svg kerak emas)."""
    urls = [m.group(2) for m in URL_RE.finditer(block)]
    for ext, fmt in FONT_RANK:
        best = next((u for u in urls if u.split("?")[0].split("#")[0].lower().endswith(ext)), None)
        if best:
            body = re.sub(r"\bsrc\s*:[^;}]*;?", "", block)
            return body.rstrip("}").rstrip() + f'\n  src: url("{best}") format("{fmt}");\n}}'
    return block
SKIP_SCHEMES = ("data:", "blob:", "javascript:", "mailto:", "tel:", "#", "about:")


class Inliner:
    def __init__(self, root: Path, remote: bool):
        self.root = root.resolve()
        self.remote = remote
        self.used: set[Path] = set()
        self.missing: list[str] = []
        self.cache: dict[str, tuple[bytes, str] | None] = {}

    # resurs: (baytlar, mime) yoki None
    def load(self, ref: str, base: str) -> tuple[bytes, str, str] | None:
        ref = ref.strip()
        if not ref or ref.startswith(SKIP_SCHEMES):
            return None
        if ref.startswith("//"):
            ref = "https:" + ref
        absolute = urllib.parse.urljoin(base, ref)
        absolute = absolute.split("#")[0]
        if absolute in self.cache:
            hit = self.cache[absolute]
            return (hit[0], hit[1], absolute) if hit else None
        data = mime = None
        if absolute.startswith("file://"):
            local = Path(urllib.parse.unquote(urllib.parse.urlparse(absolute).path))
            if local.is_file():
                data = local.read_bytes()
                mime = mimetypes.guess_type(local.name)[0] or "application/octet-stream"
                self.used.add(local.resolve())
            else:
                self.missing.append(ref)
        elif absolute.startswith(("http://", "https://")) and self.remote:
            try:
                req = urllib.request.Request(absolute, headers={
                    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
                                  "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"})
                with urllib.request.urlopen(req, timeout=20) as r:
                    data = r.read()
                    mime = (r.headers.get_content_type()
                            or mimetypes.guess_type(absolute)[0] or "application/octet-stream")
            except Exception as e:  # internet yo'q bo'lsa — havola o'z holicha qoladi
                log(f"  ogohlantirish: yuklab bo'lmadi {absolute} ({e.__class__.__name__})")
        self.cache[absolute] = (data, mime) if data is not None else None
        return (data, mime, absolute) if data is not None else None

    def data_uri(self, ref: str, base: str) -> str | None:
        got = self.load(ref, base)
        if not got:
            return None
        data, mime, absolute = got
        if mime == "text/css":
            data = self.css(data.decode("utf-8", "replace"), absolute).encode()
        return f"data:{mime};base64,{base64.b64encode(data).decode()}"

    def css(self, text: str, base: str, depth: int = 0) -> str:
        def imp(m):
            if depth > 8:
                return m.group(0)
            got = self.load(m.group(1) or m.group(2), base)
            if not got:
                return m.group(0)
            inner = self.css(got[0].decode("utf-8", "replace"), got[2], depth + 1)
            media = m.group(3).strip()
            return f"@media {media} {{\n{inner}\n}}" if media else inner

        text = IMPORT_RE.sub(imp, text)
        text = FONTFACE_RE.sub(lambda m: prune_font_face(m.group(0)), text)

        def url(m):
            uri = self.data_uri(m.group(2), base)
            return f'url("{uri}")' if uri else m.group(0)

        return URL_RE.sub(url, text)

    def srcset(self, value: str, base: str) -> str:
        parts = []
        for item in value.split(","):
            bits = item.strip().split()
            if not bits:
                continue
            uri = self.data_uri(bits[0], base)
            parts.append(" ".join([uri or bits[0]] + bits[1:]))
        return ", ".join(parts)


VFS_SHIM = """<script>/* html2one: ichki fayllar xaritasi (fetch/XHR/img uchun) */
(function(){var V=%s;
function k(u){try{if(typeof u!=='string')u=String(u&&u.url||u);if(/^(data|blob):/.test(u))return null;
var p=new URL(u,location.href);if(p.protocol!=='file:'&&p.origin!==location.origin)return null;
var rel=new URL(u,'https://v.local/').pathname.slice(1);return decodeURIComponent(rel);}catch(e){return null}}
function hit(u){var key=k(u);if(key==null)return null;if(V[key])return V[key];
for(var n in V){if(key.slice(-n.length-1)==='/'+n||key===n)return V[n];}return null}
var F=window.fetch;window.fetch=function(u,o){var d=hit(u);return F.call(this,d||u,o)};
var O=XMLHttpRequest.prototype.open;XMLHttpRequest.prototype.open=function(m,u){var d=hit(u);
arguments[1]=d||u;return O.apply(this,arguments)};
var S=Element.prototype.setAttribute;Element.prototype.setAttribute=function(n,v){
if(/^(src|href|poster)$/i.test(n)){var d=hit(v);if(d)v=d}return S.call(this,n,v)};
[HTMLImageElement,HTMLSourceElement,HTMLScriptElement,HTMLAudioElement,HTMLVideoElement].forEach(function(C){
var ds=Object.getOwnPropertyDescriptor(C.prototype,'src');if(!ds||!ds.set)return;
Object.defineProperty(C.prototype,'src',{get:ds.get,set:function(v){var d=hit(v);ds.set.call(this,d||v)},configurable:true})});
})();</script>"""

VFS_EXT = {".json", ".txt", ".csv", ".svg", ".png", ".jpg", ".jpeg", ".gif", ".webp", ".avif",
           ".mp3", ".wav", ".ogg", ".mp4", ".webm", ".html", ".htm", ".md", ".xml", ".js", ".css"}
CODE_EXT = {".js", ".css", ".html", ".htm", ".md", ".xml"}
VFS_LIMIT = 40 * 1024 * 1024


def build_single_html(entry: Path, root: Path, out: Path, remote: bool) -> list[str]:
    from bs4 import BeautifulSoup

    inl = Inliner(root, remote)
    base = entry.resolve().as_uri()
    soup = BeautifulSoup(entry.read_text("utf-8", "replace"), "html.parser")

    for tag in soup.find_all("base"):
        tag.decompose()

    for link in soup.find_all("link", href=True):
        rel = " ".join(link.get("rel", [])).lower()
        if "stylesheet" in rel:
            got = inl.load(link["href"], base)
            if got:
                style = soup.new_tag("style")
                if link.get("media"):
                    style["media"] = link["media"]
                style.string = inl.css(got[0].decode("utf-8", "replace"), got[2])
                link.replace_with(style)
        elif "icon" in rel or "apple-touch-icon" in rel:
            uri = inl.data_uri(link["href"], base)
            if uri:
                link["href"] = uri
        elif rel in ("preload", "prefetch", "modulepreload", "preconnect", "dns-prefetch", "manifest"):
            link.decompose()

    for script in soup.find_all("script", src=True):
        got = inl.load(script["src"], base)
        if not got:
            continue
        code = got[0].decode("utf-8", "replace").replace("</script", "<\\/script")
        del script["src"]
        for attr in ("integrity", "crossorigin", "async"):
            if attr in script.attrs:
                del script[attr]
        script.string = code

    for style in soup.find_all("style"):
        if style.string:
            style.string = inl.css(style.string, base)

    for tag in soup.find_all(style=True):
        tag["style"] = inl.css(tag["style"], base)

    attrs = [("img", "src"), ("source", "src"), ("video", "src"), ("video", "poster"),
             ("audio", "src"), ("track", "src"), ("input", "src"), ("embed", "src"),
             ("object", "data"), ("image", "href"), ("image", "xlink:href"), ("use", "href"),
             ("use", "xlink:href"), ("iframe", "src")]
    for name, attr in attrs:
        for tag in soup.find_all(name):
            val = tag.get(attr)
            if not val or (name == "use" and val.startswith("#")):
                continue
            if name == "use":
                # <use href="icons.svg#id"> — data URI ishlamaydi, faqat fayl qismini ichiga qo'yamiz
                file_part, _, frag = val.partition("#")
                got = inl.load(file_part, base) if file_part else None
                if got and frag:
                    sprite = BeautifulSoup(got[0].decode("utf-8", "replace"), "html.parser")
                    sym = sprite.find(id=frag)
                    if sym:
                        holder = soup.new_tag("svg", style="display:none")
                        holder.append(sym)
                        (soup.body or soup).insert(0, holder)
                        tag[attr] = "#" + frag
                continue
            uri = inl.data_uri(val, base)
            if uri:
                tag[attr] = uri
    for name in ("img", "source"):
        for tag in soup.find_all(name, srcset=True):
            tag["srcset"] = inl.srcset(tag["srcset"], base)
    for tag in soup.find_all(attrs={"data-src": True}):
        uri = inl.data_uri(tag["data-src"], base)
        if uri:
            tag["data-src"] = uri
    for tag in soup.find_all(attrs={"data-background-image": True}):  # reveal.js
        uri = inl.data_uri(tag["data-background-image"], base)
        if uri:
            tag["data-background-image"] = uri

    # JS ichidan dinamik chaqiriladigan fayllar — xarita orqali. HTML/CSS'da ichiga qo'yilgan
    # fayl JS kodida nomi bilan tilga olingan bo'lsa, u ham xaritaga qo'shiladi.
    js_text = "\n".join(s.string or "" for s in soup.find_all("script"))
    vfs, total = {}, 0
    for f in sorted(root.rglob("*")):
        if not f.is_file() or f.resolve() == entry.resolve():
            continue
        named_in_js = f.name in js_text
        is_code = f.suffix.lower() in CODE_EXT
        if f.resolve() in inl.used and (is_code or not named_in_js):
            continue
        # ishlatilmagan kod/sahifa fayllari (boshqa mavzular, pluginlar) faqat JS'da nomi bo'lsa
        if is_code and not named_in_js:
            continue
        if f.suffix.lower() not in VFS_EXT or total + f.stat().st_size > VFS_LIMIT:
            continue
        mime = mimetypes.guess_type(f.name)[0] or "application/octet-stream"
        rel = f.relative_to(root).as_posix()
        vfs[rel] = f"data:{mime};base64,{base64.b64encode(f.read_bytes()).decode()}"
        total += f.stat().st_size
    if vfs:
        shim = BeautifulSoup(VFS_SHIM % json.dumps(vfs), "html.parser")
        head = soup.head or soup
        head.insert(0, shim)

    if soup.head and not soup.head.find("meta", charset=True):
        soup.head.insert(0, soup.new_tag("meta", charset="utf-8"))

    out.write_text(str(soup), "utf-8")
    return inl.missing


# ---------------------------------------------------------------- main

def main() -> None:
    ap = argparse.ArgumentParser(description="Ko'p faylli HTML → bitta PPTX / PDF / HTML")
    ap.add_argument("source", nargs="?", default="~/Downloads/anor-delivery.zip",
                    help="zip, papka yoki index.html (standart: ~/Downloads/anor-delivery.zip)")
    ap.add_argument("--out", help="natijalar papkasi")
    ap.add_argument("--name", help="natija fayllar nomi (standart: zip nomi)")
    ap.add_argument("--size", default="1920x1080")
    ap.add_argument("--mode", default="auto", choices=["auto", "reveal", "sections", "keys", "page"])
    ap.add_argument("--selector")
    ap.add_argument("--wait", type=int, default=800)
    ap.add_argument("--scale", type=float, default=2)
    ap.add_argument("--png", action="store_true", help="JPEG o'rniga PNG (sifat yuqori, fayl katta)")
    ap.add_argument("--only", default="pptx,pdf,html")
    ap.add_argument("--no-remote", action="store_true")
    args = ap.parse_args()

    src = Path(args.source).expanduser().resolve()
    if not src.exists():
        sys.exit(f"Xato: {src} topilmadi")
    args.width, args.height = (int(x) for x in args.size.lower().split("x"))
    want = {x.strip() for x in args.only.split(",")}
    name = args.name or (src.stem if src.suffix.lower() != ".html" else src.parent.name)
    out_dir = Path(args.out).expanduser().resolve() if args.out else src.parent
    out_dir.mkdir(parents=True, exist_ok=True)

    work = Path(tempfile.mkdtemp(prefix="html2one-"))
    try:
        root = prepare_source(src, work)
        entry = entry_file(root, src)
        log(f"Manba: {entry}")
        results = []

        if want & {"pptx", "pdf"}:
            httpd, port = serve(root)
            try:
                rel = urllib.parse.quote(entry.relative_to(root).as_posix())
                log("Slaydlar suratga olinmoqda...")
                slides = capture(f"http://127.0.0.1:{port}/{rel}", args)
            finally:
                httpd.shutdown()
            log(f"  {len(slides)} ta slayd")
            if not slides:
                sys.exit("Xato: sahifadan hech narsa olinmadi")
            if "pptx" in want:
                p = out_dir / f"{name}.pptx"
                write_pptx(slides, p, args.width, args.height)
                results.append(p)
            if "pdf" in want:
                p = out_dir / f"{name}.pdf"
                write_pdf(slides, p, args.width)
                results.append(p)

        if "html" in want:
            p = out_dir / f"{name}.html"
            if p.resolve() == entry.resolve():
                p = out_dir / f"{name}.single.html"
            missing = build_single_html(entry, root, p, remote=not args.no_remote)
            for m in sorted(set(missing)):
                log(f"  ogohlantirish: fayl topilmadi — {m}")
            results.append(p)

        log("\nTayyor:")
        for p in results:
            log(f"  {p}  ({p.stat().st_size / 1024 / 1024:.1f} MB)")
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main()
